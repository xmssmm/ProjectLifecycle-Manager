from __future__ import annotations

from collections.abc import AsyncIterator, Callable
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from time import perf_counter
from typing import cast
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.api.v1.reports import get_report_service
from app.core.db import get_db_session
from app.core.deps import get_current_user
from app.core.exceptions import PermissionDeniedError
from app.core.middleware import InMemoryRateLimitStore
from app.main import create_app
from app.models.departments import Department
from app.models.main_projects import MainProject, MainProjectStatus
from app.models.payments import Payment, PaymentType
from app.models.reports import ReportJob, ReportJobStatus, ReportType
from app.models.sub_projects import SubProject
from app.models.users import User, UserRole
from app.schemas.reports import ReportCreate
from app.services.reports import (
    InMemoryReportRepository,
    ReportDownload,
    ReportFileFormat,
    ReportService,
    ReportTaskDispatcher,
)
from app.storage.base import StorageBackend
from tests.factories import (
    DepartmentFactory,
    MainProjectFactory,
    PaymentFactory,
    SubProjectFactory,
    UserFactory,
)


@dataclass(frozen=True)
class SavedReport:
    storage_key: str
    filename: str
    content: bytes


class RecordingStorage(StorageBackend):
    def __init__(self) -> None:
        self.saved: list[SavedReport] = []

    def save(
        self,
        *,
        sub_id: str,
        phase_id: str,
        filename: str,
        content: bytes,
    ) -> str:
        storage_key = f"{sub_id}/{phase_id}/{filename}"
        self.saved.append(SavedReport(storage_key=storage_key, filename=filename, content=content))
        return storage_key

    def read(self, storage_key: str) -> bytes:
        return next(item.content for item in self.saved if item.storage_key == storage_key)

    def delete(self, storage_key: str) -> None:
        self.saved = [item for item in self.saved if item.storage_key != storage_key]

    def get_url(self, storage_key: str) -> str:
        return f"/storage/{storage_key}"


class RecordingDispatcher(ReportTaskDispatcher):
    def __init__(self) -> None:
        self.enqueued: list[UUID] = []

    def enqueue(self, job_id: UUID) -> None:
        self.enqueued.append(job_id)


@pytest.mark.asyncio
async def test_report_service_generates_all_templates_as_excel_and_pdf() -> None:
    actor = cast(User, UserFactory(role=UserRole.dept_manager))
    service, repository, storage, _dispatcher = make_service(async_threshold_rows=1000)

    for report_type in ReportType:
        job = await service.create_report(
            actor=actor,
            payload=ReportCreate(report_type=report_type, parameters={}),
        )

        assert job.status == ReportJobStatus.completed
        assert job.progress == 100
        assert job.row_count > 0
        assert job.xlsx_storage_key is not None
        assert job.pdf_storage_key is not None
        assert repository.jobs[job.id] == job

    assert len(storage.saved) == len(ReportType) * 2
    xlsx_files = [item for item in storage.saved if item.filename.endswith(".xlsx")]
    pdf_files = [item for item in storage.saved if item.filename.endswith(".pdf")]
    assert len(xlsx_files) == len(ReportType)
    assert len(pdf_files) == len(ReportType)
    assert all(load_workbook(BytesIO(item.content)).active.max_row >= 2 for item in xlsx_files)
    assert all(item.content.startswith(b"%PDF") for item in pdf_files)


@pytest.mark.asyncio
async def test_large_report_is_queued_and_worker_updates_progress() -> None:
    actor = cast(User, UserFactory(role=UserRole.dept_manager))
    service, repository, storage, dispatcher = make_service(
        sub_project_count=1000,
        async_threshold_rows=1000,
    )

    job = await service.create_report(
        actor=actor,
        payload=ReportCreate(report_type=ReportType.project_list, parameters={}),
    )

    assert job.status == ReportJobStatus.queued
    assert job.progress == 0
    assert job.row_count == 1000
    assert dispatcher.enqueued == [job.id]
    assert storage.saved == []

    started_at = perf_counter()
    completed = await service.run_report_job(job.id)
    elapsed_seconds = perf_counter() - started_at

    assert completed.status == ReportJobStatus.completed
    assert completed.progress == 100
    assert completed.xlsx_storage_key is not None
    assert completed.pdf_storage_key is not None
    assert len(storage.saved) == 2
    assert repository.jobs[job.id].status == ReportJobStatus.completed
    assert elapsed_seconds < 30


@pytest.mark.asyncio
async def test_report_service_enforces_generator_role() -> None:
    actor = cast(User, UserFactory(role=UserRole.proj_member))
    service, _repository, _storage, _dispatcher = make_service()

    with pytest.raises(PermissionDeniedError):
        await service.create_report(
            actor=actor,
            payload=ReportCreate(report_type=ReportType.project_list, parameters={}),
        )


def test_report_endpoints_create_job_and_return_progress_payload() -> None:
    requester = cast(User, UserFactory(role=UserRole.dept_manager))
    job = ReportJob(
        id=uuid4(),
        report_type=ReportType.project_list,
        requested_by_id=requester.id,
        parameters={},
        status=ReportJobStatus.completed,
        progress=100,
        row_count=2,
        xlsx_storage_key="reports/job/report.xlsx",
        pdf_storage_key="reports/job/report.pdf",
        error_message=None,
        started_at=None,
        finished_at=None,
    )

    class FakeReportService:
        async def create_report(self, *, actor: User, payload: ReportCreate) -> ReportJob:
            assert actor.id == requester.id
            assert payload.report_type == ReportType.project_list
            return job

        async def get_report_job(self, *, actor: User, job_id: UUID) -> ReportJob:
            assert actor.id == requester.id
            assert job_id == job.id
            return job

    async def fake_db_session() -> AsyncIterator[object]:
        yield object()

    async def fake_current_user() -> User:
        return requester

    async def fake_report_service() -> FakeReportService:
        return FakeReportService()

    app = create_app(rate_limit_store=InMemoryRateLimitStore())
    app.dependency_overrides[get_db_session] = fake_db_session
    app.dependency_overrides[get_current_user] = fake_current_user
    app.dependency_overrides[get_report_service] = fake_report_service
    client = TestClient(app)

    create_response = client.post(
        "/api/v1/reports",
        json={"report_type": "project_list", "parameters": {}},
    )
    progress_response = client.get(f"/api/v1/reports/{job.id}")

    assert create_response.status_code == 200
    assert create_response.json()["data"]["status"] == "completed"
    assert create_response.json()["data"]["available_formats"] == ["xlsx", "pdf"]
    assert progress_response.status_code == 200
    payload = progress_response.json()["data"]
    assert payload["id"] == str(job.id)
    assert payload["report_type"] == "project_list"
    assert payload["progress"] == 100
    assert payload["row_count"] == 2


@pytest.mark.asyncio
async def test_report_download_returns_requested_format_for_requester_only() -> None:
    requester = cast(User, UserFactory(role=UserRole.dept_manager))
    other_user = cast(User, UserFactory(role=UserRole.dept_manager))
    service, repository, storage, _dispatcher = make_service()
    job = make_completed_report_job(requester)
    repository.jobs[job.id] = job
    storage.saved = [
        SavedReport(job.xlsx_storage_key or "", "project_list.xlsx", b"xlsx-bytes"),
        SavedReport(job.pdf_storage_key or "", "project_list.pdf", b"%PDF-1.7"),
    ]

    download = await service.download_report(
        actor=requester,
        job_id=job.id,
        file_format=ReportFileFormat.xlsx,
    )

    assert download.file_name == "project_list.xlsx"
    assert (
        download.content_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert download.content == b"xlsx-bytes"
    with pytest.raises(PermissionDeniedError):
        await service.download_report(
            actor=other_user,
            job_id=job.id,
            file_format=ReportFileFormat.pdf,
        )


def test_report_download_endpoint_streams_file_with_content_disposition() -> None:
    requester = cast(User, UserFactory(role=UserRole.dept_manager))
    job_id = uuid4()

    class FakeReportService:
        async def download_report(
            self,
            *,
            actor: User,
            job_id: UUID,
            file_format: ReportFileFormat,
        ) -> ReportDownload:
            assert actor.id == requester.id
            assert file_format == ReportFileFormat.xlsx
            return ReportDownload(
                file_name="project_list.xlsx",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                content=b"xlsx-bytes",
            )

    async def fake_db_session() -> AsyncIterator[object]:
        yield object()

    async def fake_current_user() -> User:
        return requester

    async def fake_report_service() -> FakeReportService:
        return FakeReportService()

    app = create_app(rate_limit_store=InMemoryRateLimitStore())
    app.dependency_overrides[get_db_session] = fake_db_session
    app.dependency_overrides[get_current_user] = fake_current_user
    app.dependency_overrides[get_report_service] = fake_report_service

    response = TestClient(app).get(f"/api/v1/reports/{job_id}/download?format=xlsx")

    assert response.status_code == 200
    assert response.content == b"xlsx-bytes"
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert "project_list.xlsx" in response.headers["content-disposition"]


@pytest.mark.asyncio
async def test_report_cleanup_deletes_files_after_retention_window() -> None:
    now = datetime(2026, 5, 10, tzinfo=UTC)
    requester = cast(User, UserFactory(role=UserRole.dept_manager))
    service, repository, storage, _dispatcher = make_service(now_provider=lambda: now)
    expired_job = make_completed_report_job(
        requester,
        finished_at=now - timedelta(days=8),
    )
    fresh_job = make_completed_report_job(
        requester,
        finished_at=now - timedelta(days=6),
    )
    repository.jobs[expired_job.id] = expired_job
    repository.jobs[fresh_job.id] = fresh_job
    storage.saved = [
        SavedReport(expired_job.xlsx_storage_key or "", "old.xlsx", b"old-xlsx"),
        SavedReport(expired_job.pdf_storage_key or "", "old.pdf", b"old-pdf"),
        SavedReport(fresh_job.xlsx_storage_key or "", "fresh.xlsx", b"fresh-xlsx"),
        SavedReport(fresh_job.pdf_storage_key or "", "fresh.pdf", b"fresh-pdf"),
    ]

    result = await service.cleanup_expired_reports(retention_days=7)

    assert result.deleted_jobs == 1
    assert result.deleted_files == 2
    assert expired_job.xlsx_storage_key is None
    assert expired_job.pdf_storage_key is None
    assert fresh_job.xlsx_storage_key is not None
    assert fresh_job.pdf_storage_key is not None
    assert [item.filename for item in storage.saved] == ["fresh.xlsx", "fresh.pdf"]


def test_celery_beat_schedules_report_cleanup_daily() -> None:
    from app.core.config import Settings
    from app.tasks.celery_app import create_celery_app
    from app.tasks.task_names import REPORT_CLEANUP_TASK_NAME

    celery_app = create_celery_app(Settings(redis_url="redis://redis:6379/0"))

    schedule = celery_app.conf.beat_schedule["report-cleanup-daily-0330"]
    assert schedule["task"] == REPORT_CLEANUP_TASK_NAME


def make_service(
    *,
    sub_project_count: int = 2,
    async_threshold_rows: int = 1000,
    now_provider: Callable[[], datetime] | None = None,
) -> tuple[ReportService, InMemoryReportRepository, RecordingStorage, RecordingDispatcher]:
    department = cast(Department, DepartmentFactory(code="general", name="综合部"))
    main_project = cast(
        MainProject,
        MainProjectFactory(
            dept_id=department.id,
            project_no="Z-2026-0001",
            name="智慧档案平台",
            status=MainProjectStatus.in_progress,
            total_budget=Decimal("500000.00"),
            spent_amount=Decimal("30000.00"),
        ),
    )
    sub_projects = [
        cast(
            SubProject,
            SubProjectFactory(
                main_project_id=main_project.id,
                dept_id=department.id,
                project_no=f"Z-2026-0001-ZX-{index + 1:03d}",
                name=f"子项目 {index + 1}",
                budget=Decimal("1000.00"),
                spent_amount=Decimal("100.00"),
            ),
        )
        for index in range(sub_project_count)
    ]
    payments = [
        cast(
            Payment,
            PaymentFactory(
                sub_project_id=sub_projects[0].id,
                payment_no="PAY-2026-0001",
                amount=Decimal("100.00"),
                payment_date=date(2026, 5, 10),
                payment_type=PaymentType.normal,
            ),
        ),
    ]
    repository = InMemoryReportRepository(
        departments=[department],
        main_projects=[main_project],
        sub_projects=sub_projects,
        payments=payments,
    )
    storage = RecordingStorage()
    dispatcher = RecordingDispatcher()
    service = ReportService(
        repository=repository,
        storage=storage,
        dispatcher=dispatcher,
        async_threshold_rows=async_threshold_rows,
        now_provider=now_provider if now_provider is not None else lambda: datetime.now(UTC),
    )
    return service, repository, storage, dispatcher


def make_completed_report_job(
    requester: User,
    *,
    finished_at: datetime | None = None,
) -> ReportJob:
    now = finished_at or datetime(2026, 5, 10, tzinfo=UTC)
    job_id = uuid4()
    return ReportJob(
        id=job_id,
        report_type=ReportType.project_list,
        requested_by_id=requester.id,
        parameters={},
        status=ReportJobStatus.completed,
        progress=100,
        row_count=2,
        xlsx_storage_key=f"reports/{job_id}/project_list.xlsx",
        pdf_storage_key=f"reports/{job_id}/project_list.pdf",
        error_message=None,
        started_at=now,
        finished_at=now,
        created_at=now,
        updated_at=now,
    )
