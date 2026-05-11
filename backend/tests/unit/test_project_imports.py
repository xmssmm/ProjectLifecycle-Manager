from __future__ import annotations

from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from app.models.departments import Department
from app.models.main_projects import MainProject, MainProjectStatus
from app.models.users import User, UserRole, UserStatus
from app.services.audit import InMemoryAuditLogWriter
from app.services.project_imports import (
    PROJECT_IMPORT_TEMPLATE_HEADERS,
    InMemoryProjectImportRepository,
    ProjectImportService,
)

NOW = datetime(2026, 5, 11, 12, 0, tzinfo=UTC)


def make_user(role: UserRole = UserRole.admin) -> User:
    return User(
        id=uuid4(),
        username=f"{role.value}-{uuid4().hex[:6]}",
        email=None,
        password_hash="hashed",
        role=role,
        dept_id=None,
        status=UserStatus.active,
        password_changed_at=NOW,
        last_login_at=None,
        created_at=NOW,
        updated_at=NOW,
    )


def make_department(*, code: str = "D001") -> Department:
    return Department(
        id=uuid4(),
        code=code,
        name=f"Department {code}",
        created_at=NOW,
        updated_at=NOW,
    )


def make_main_project(*, project_no: str, dept_id: UUID | None = None) -> MainProject:
    return MainProject(
        id=uuid4(),
        project_no=project_no,
        name=f"Existing {project_no}",
        dept_id=uuid4() if dept_id is None else dept_id,
        status=MainProjectStatus.not_started,
        total_budget=Decimal("100000.00"),
        expected_finish_date=date(2026, 12, 31),
        spent_amount=Decimal("0.00"),
        remark=None,
        creator_id=uuid4(),
        created_at=NOW,
        updated_at=NOW,
    )


def build_workbook(rows: list[list[object | None]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PROJECT_IMPORT_TEMPLATE_HEADERS)
    for row in rows:
        worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_service(
    *,
    departments: list[Department],
    projects: list[MainProject] | None = None,
) -> tuple[ProjectImportService, InMemoryProjectImportRepository]:
    repository = InMemoryProjectImportRepository(
        departments=departments,
        projects=projects or [],
    )
    return (
        ProjectImportService(
            repository=repository,
            now_provider=lambda: NOW,
            batch_token_provider=lambda: "ABCDEF12",
        ),
        repository,
    )


def test_project_import_template_contains_complete_headers() -> None:
    department = make_department()
    service, _repository = make_service(departments=[department])

    content = service.generate_template()
    worksheet = load_workbook(BytesIO(content)).active

    assert [cell.value for cell in worksheet[1]] == [
        "项目编号",
        "项目名称",
        "部门编码",
        "总预算",
        "预计完成日期",
        "备注",
    ]


@pytest.mark.asyncio
async def test_project_import_allows_partial_failures_and_records_audit() -> None:
    department = make_department(code="D001")
    existing = make_main_project(project_no="Z-2026-0002", dept_id=department.id)
    service, repository = make_service(departments=[department], projects=[existing])
    audit_writer = InMemoryAuditLogWriter()
    content = build_workbook(
        [
            ["Z-2026-1001", "Imported A", "D001", "10000.50", "2026-12-31", "first"],
            ["Z-2026-0002", "Duplicate Existing", "D001", "9000", "2026-12-31", None],
            ["Z-2026-1003", "Missing Department", "NOPE", "9000", "2026-12-31", None],
            ["Z-2026-1004", "Imported B", "D001", Decimal("12000.00"), date(2026, 11, 30), None],
        ],
    )

    result = await service.import_projects(
        actor=make_user(),
        content=content,
        audit_writer=audit_writer,
    )

    assert result.batch_no == "IMPORT-20260511-ABCDEF12"
    assert result.total_rows == 4
    assert result.success_count == 2
    assert result.failure_count == 2
    assert [error.row_number for error in result.errors] == [3, 4]
    assert [error.field for error in result.errors] == ["project_no", "dept_code"]
    assert [project.project_no for project in repository.projects] == [
        "Z-2026-0002",
        "Z-2026-1001",
        "Z-2026-1004",
    ]
    assert repository.projects[-1].expected_finish_date == date(2026, 11, 30)
    assert repository.projects[-1].status == MainProjectStatus.not_started
    assert audit_writer.entries[0].action == "project_import.create_batch"
    assert audit_writer.entries[0].target_id == "IMPORT-20260511-ABCDEF12"
    assert audit_writer.entries[0].after_state["success_count"] == 2


@pytest.mark.asyncio
async def test_project_import_rejects_duplicate_project_no_inside_workbook() -> None:
    department = make_department(code="D001")
    service, repository = make_service(departments=[department])
    content = build_workbook(
        [
            ["Z-2026-2001", "Imported A", "D001", "10000", "2026-12-31", None],
            ["Z-2026-2001", "Duplicate A", "D001", "12000", "2026-12-31", None],
        ],
    )

    result = await service.import_projects(actor=make_user(), content=content)

    assert result.success_count == 1
    assert result.failure_count == 1
    assert result.errors[0].row_number == 3
    assert result.errors[0].field == "project_no"
    assert [project.project_no for project in repository.projects] == ["Z-2026-2001"]


@pytest.mark.asyncio
async def test_project_import_records_1000_row_performance_baseline() -> None:
    department = make_department(code="D001")
    service, repository = make_service(departments=[department])
    rows: list[list[object | None]] = [
        [f"Z-2026-{index:04d}", f"Imported {index}", "D001", "1000.00", "2026-12-31", None]
        for index in range(1, 1001)
    ]

    result = await service.import_projects(actor=make_user(), content=build_workbook(rows))

    assert result.total_rows == 1000
    assert result.success_count == 1000
    assert result.failure_count == 0
    assert result.duration_ms < 60000
    assert len(repository.projects) == 1000
