from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.models.exports import DatabaseExportJobStatus
from app.models.users import User, UserRole, UserStatus
from app.services.audit import InMemoryAuditLogWriter
from app.services.database_exports import (
    DatabaseExportForeignKey,
    DatabaseExportService,
    DatabaseExportTable,
    InMemoryDatabaseExportRepository,
)
from app.storage.base import StorageBackend

NOW = datetime(2026, 5, 11, 12, 0, tzinfo=UTC)


@dataclass(frozen=True)
class SavedFile:
    storage_key: str
    filename: str
    content: bytes


class RecordingStorage(StorageBackend):
    def __init__(self) -> None:
        self.saved: list[SavedFile] = []

    def save(self, *, sub_id: str, phase_id: str, filename: str, content: bytes) -> str:
        storage_key = f"{sub_id}/{phase_id}/{filename}"
        self.saved.append(SavedFile(storage_key=storage_key, filename=filename, content=content))
        return storage_key

    def read(self, storage_key: str) -> bytes:
        return next(item.content for item in self.saved if item.storage_key == storage_key)

    def delete(self, storage_key: str) -> None:
        self.saved = [item for item in self.saved if item.storage_key != storage_key]

    def get_url(self, storage_key: str) -> str:
        return f"/storage/{storage_key}"


class RecordingDispatcher:
    def __init__(self) -> None:
        self.enqueued_job_ids: list[str] = []

    def enqueue(self, job_id: object) -> None:
        self.enqueued_job_ids.append(str(job_id))


def make_admin() -> User:
    return User(
        id=uuid4(),
        username="admin",
        email=None,
        password_hash="hashed",
        role=UserRole.admin,
        dept_id=None,
        status=UserStatus.active,
        password_changed_at=NOW,
        last_login_at=None,
        created_at=NOW,
        updated_at=NOW,
    )


def make_service() -> tuple[
    DatabaseExportService,
    InMemoryDatabaseExportRepository,
    RecordingStorage,
    RecordingDispatcher,
]:
    tables = [
        DatabaseExportTable(
            name="departments",
            columns=["id", "code", "name"],
            rows=[
                {
                    "id": "dept-1",
                    "code": "D001",
                    "name": "Digital",
                },
            ],
        ),
        DatabaseExportTable(
            name="main_projects",
            columns=["id", "project_no", "dept_id"],
            rows=[
                {
                    "id": "project-1",
                    "project_no": "Z-2026-0001",
                    "dept_id": "dept-1",
                },
            ],
            foreign_keys=[
                DatabaseExportForeignKey(
                    table="main_projects",
                    column="dept_id",
                    referred_table="departments",
                    referred_column="id",
                ),
            ],
        ),
    ]
    repository = InMemoryDatabaseExportRepository(tables=tables)
    storage = RecordingStorage()
    dispatcher = RecordingDispatcher()
    return (
        DatabaseExportService(
            repository=repository,
            storage=storage,
            dispatcher=dispatcher,
            now_provider=lambda: NOW,
        ),
        repository,
        storage,
        dispatcher,
    )


@pytest.mark.asyncio
async def test_create_database_export_job_enqueues_async_work_and_audits() -> None:
    service, repository, _storage, dispatcher = make_service()
    admin = make_admin()
    audit_writer = InMemoryAuditLogWriter()

    job = await service.create_export(actor=admin, audit_writer=audit_writer)

    assert job.status == DatabaseExportJobStatus.queued
    assert repository.jobs[job.id] == job
    assert dispatcher.enqueued_job_ids == [str(job.id)]
    assert audit_writer.entries[0].action == "database_export.create_job"
    assert audit_writer.entries[0].target_id == str(job.id)


@pytest.mark.asyncio
async def test_database_export_package_contains_manifest_csv_xlsx_and_fk_references() -> None:
    service, repository, storage, _dispatcher = make_service()
    admin = make_admin()
    job = await service.create_export(actor=admin)

    completed = await service.run_export_job(job.id)

    assert completed.status == DatabaseExportJobStatus.completed
    assert completed.progress == 100
    assert completed.table_count == 2
    assert completed.row_count == 2
    assert completed.storage_key is not None
    package = storage.read(completed.storage_key)
    with ZipFile(BytesIO(package)) as archive:
        assert {
            "manifest.json",
            "foreign_keys.json",
            "audit.json",
            "csv/departments.csv",
            "csv/main_projects.csv",
            "xlsx/database_export.xlsx",
        }.issubset(set(archive.namelist()))
        manifest = json.loads(archive.read("manifest.json"))
        foreign_keys = json.loads(archive.read("foreign_keys.json"))
        csv_content = archive.read("csv/main_projects.csv").decode("utf-8-sig")
        workbook = load_workbook(BytesIO(archive.read("xlsx/database_export.xlsx")))

    assert manifest["generated_at"] == "2026-05-11T12:00:00+00:00"
    assert manifest["operator_id"] == str(admin.id)
    assert manifest["tables"][0] == {
        "name": "departments",
        "row_count": 1,
        "columns": ["id", "code", "name"],
    }
    assert foreign_keys == [
        {
            "table": "main_projects",
            "column": "dept_id",
            "referred_table": "departments",
            "referred_column": "id",
        },
    ]
    assert "Z-2026-0001" in csv_content
    assert "main_projects" in workbook.sheetnames
    assert repository.jobs[job.id].manifest["total_rows"] == 2


@pytest.mark.asyncio
async def test_database_export_download_reads_zip_and_records_audit() -> None:
    service, _repository, _storage, _dispatcher = make_service()
    admin = make_admin()
    job = await service.create_export(actor=admin)
    completed = await service.run_export_job(job.id)
    audit_writer = InMemoryAuditLogWriter()

    download = await service.download_export(
        actor=admin,
        job_id=completed.id,
        audit_writer=audit_writer,
    )

    assert download.file_name == "database_export.zip"
    assert download.content_type == "application/zip"
    assert download.content.startswith(b"PK")
    assert audit_writer.entries[0].action == "database_export.download"
    assert audit_writer.entries[0].target_id == str(completed.id)
