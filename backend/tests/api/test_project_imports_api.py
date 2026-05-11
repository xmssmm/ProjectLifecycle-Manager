from __future__ import annotations

from collections.abc import AsyncIterator
from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from app.api.v1.imports import get_project_import_service
from app.core.db import get_db_session
from app.core.deps import get_current_user
from app.core.middleware import InMemoryRateLimitStore
from app.main import create_app
from app.models.main_projects import MainProjectStatus
from app.models.users import User, UserRole, UserStatus
from app.services.project_imports import (
    PROJECT_IMPORT_TEMPLATE_HEADERS,
    ProjectImportCreatedProject,
    ProjectImportResult,
    ProjectImportRowError,
)

NOW = datetime(2026, 5, 11, 12, 0, tzinfo=UTC)


def make_user(role: UserRole) -> User:
    return User(
        id=uuid4(),
        username=f"{role.value}-{uuid4().hex[:6]}",
        email=None,
        password_hash="hashed",
        role=role,
        dept_id=None,
        status=UserStatus.active,
        password_changed_at=NOW,
        last_login_at=None,
        created_at=NOW,
        updated_at=NOW,
    )


def build_template() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PROJECT_IMPORT_TEMPLATE_HEADERS)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_client(user: User, service: object) -> TestClient:
    async def fake_db_session() -> AsyncIterator[object]:
        yield object()

    async def fake_current_user() -> User:
        return user

    async def fake_service() -> object:
        return service

    app = create_app(rate_limit_store=InMemoryRateLimitStore())
    app.dependency_overrides[get_db_session] = fake_db_session
    app.dependency_overrides[get_current_user] = fake_current_user
    app.dependency_overrides[get_project_import_service] = fake_service
    return TestClient(app)


def test_project_import_api_downloads_template_for_admin() -> None:
    class FakeProjectImportService:
        def generate_template(self) -> bytes:
            return build_template()

    client = build_client(make_user(UserRole.admin), FakeProjectImportService())

    response = client.get("/api/v1/imports/projects/template")

    assert response.status_code == 200
    assert "project_import_template.xlsx" in response.headers["content-disposition"]
    worksheet = load_workbook(BytesIO(response.content)).active
    assert [cell.value for cell in worksheet[1]] == PROJECT_IMPORT_TEMPLATE_HEADERS


def test_project_import_api_uploads_workbook_and_returns_row_errors() -> None:
    admin = make_user(UserRole.admin)

    class FakeProjectImportService:
        async def import_projects(
            self,
            *,
            actor: User,
            content: bytes,
            audit_writer: object | None = None,
            audit_context: object | None = None,
        ) -> ProjectImportResult:
            assert actor.id == admin.id
            assert content.startswith(b"PK")
            assert audit_writer is not None
            assert audit_context is not None
            return ProjectImportResult(
                batch_no="IMPORT-20260511-ABCDEF12",
                total_rows=2,
                success_count=1,
                failure_count=1,
                duration_ms=12,
                errors=[
                    ProjectImportRowError(
                        row_number=3,
                        field="project_no",
                        message="项目编号重复",
                        value="Z-2026-0001",
                    ),
                ],
                created_projects=[
                    ProjectImportCreatedProject(
                        id=uuid4(),
                        project_no="Z-2026-1001",
                        name="Imported",
                        dept_id=uuid4(),
                        status=MainProjectStatus.not_started,
                    ),
                ],
            )

    client = build_client(admin, FakeProjectImportService())

    response = client.post(
        "/api/v1/imports/projects",
        files={
            "file": (
                "projects.xlsx",
                build_template(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )

    assert response.status_code == 200
    assert response.json()["data"]["batch_no"] == "IMPORT-20260511-ABCDEF12"
    assert response.json()["data"]["success_count"] == 1
    assert response.json()["data"]["errors"][0]["row_number"] == 3


def test_project_import_api_rejects_non_admin_users() -> None:
    client = build_client(make_user(UserRole.dept_manager), object())

    template_response = client.get("/api/v1/imports/projects/template")
    import_response = client.post(
        "/api/v1/imports/projects",
        files={"file": ("projects.xlsx", build_template())},
    )

    assert template_response.status_code == 403
    assert import_response.status_code == 403
