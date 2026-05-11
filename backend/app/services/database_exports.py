from __future__ import annotations

import csv
import json
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Any, Protocol, cast
from uuid import UUID, uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook  # type: ignore[import-untyped]
from sqlalchemy import Table, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import BusinessException, PermissionDeniedError, ResourceNotFoundError
from app.models import Base
from app.models.exports import DatabaseExportJob, DatabaseExportJobStatus
from app.models.users import User, UserRole
from app.services.audit import AuditContext, AuditLogEntry, AuditLogWriter
from app.storage.base import StorageBackend


@dataclass(frozen=True)
class DatabaseExportForeignKey:
    table: str
    column: str
    referred_table: str
    referred_column: str


@dataclass(frozen=True)
class DatabaseExportTable:
    name: str
    columns: list[str]
    rows: list[dict[str, object]]
    foreign_keys: list[DatabaseExportForeignKey] = field(default_factory=list)


@dataclass(frozen=True)
class DatabaseExportDownload:
    file_name: str
    content_type: str
    content: bytes


@dataclass(frozen=True)
class TabularExportColumn:
    key: str
    label: str


class TabularExportRenderer:
    @staticmethod
    def render_csv(
        *,
        columns: Sequence[TabularExportColumn],
        rows: Sequence[Mapping[str, object]],
    ) -> bytes:
        output = StringIO()
        writer = csv.DictWriter(
            output,
            fieldnames=[column.key for column in columns],
            lineterminator="\n",
        )
        writer.writerow({column.key: column.label for column in columns})
        for row in rows:
            writer.writerow(
                {
                    column.key: DatabaseExportService.csv_value(row.get(column.key))
                    for column in columns
                },
            )
        return output.getvalue().encode("utf-8-sig")

    @staticmethod
    def render_xlsx(
        *,
        columns: Sequence[TabularExportColumn],
        rows: Sequence[Mapping[str, object]],
        summary: Mapping[str, object],
    ) -> bytes:
        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "summary"
        for key, value in summary.items():
            summary_sheet.append([key, DatabaseExportService.csv_value(value)])

        data_sheet = workbook.create_sheet("data")
        data_sheet.append([column.key for column in columns])
        for row in rows:
            data_sheet.append(
                [DatabaseExportService.csv_value(row.get(column.key)) for column in columns],
            )
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def render_zip(
        *,
        columns: Sequence[TabularExportColumn],
        rows: Sequence[Mapping[str, object]],
        summary: Mapping[str, object],
    ) -> bytes:
        buffer = BytesIO()
        with ZipFile(buffer, mode="w", compression=ZIP_DEFLATED) as archive:
            archive.writestr(
                "custom_report.csv",
                TabularExportRenderer.render_csv(columns=columns, rows=rows),
            )
            archive.writestr(
                "custom_report.xlsx",
                TabularExportRenderer.render_xlsx(
                    columns=columns,
                    rows=rows,
                    summary=summary,
                ),
            )
        return buffer.getvalue()


class DatabaseExportTaskDispatcher(Protocol):
    def enqueue(self, job_id: UUID) -> None:
        ...


class NoopDatabaseExportTaskDispatcher:
    def enqueue(self, job_id: UUID) -> None:
        _ = job_id


class DatabaseExportRepository(Protocol):
    async def add_job(self, job: DatabaseExportJob) -> DatabaseExportJob:
        ...

    async def get_job(self, job_id: UUID) -> DatabaseExportJob | None:
        ...

    async def save_job(self, job: DatabaseExportJob) -> DatabaseExportJob:
        ...

    async def load_tables(self) -> list[DatabaseExportTable]:
        ...


class InMemoryDatabaseExportRepository:
    def __init__(self, *, tables: Sequence[DatabaseExportTable] | None = None) -> None:
        self.tables = list(tables or [])
        self.jobs: dict[UUID, DatabaseExportJob] = {}

    async def add_job(self, job: DatabaseExportJob) -> DatabaseExportJob:
        self.jobs[job.id] = job
        return job

    async def get_job(self, job_id: UUID) -> DatabaseExportJob | None:
        return self.jobs.get(job_id)

    async def save_job(self, job: DatabaseExportJob) -> DatabaseExportJob:
        self.jobs[job.id] = job
        return job

    async def load_tables(self) -> list[DatabaseExportTable]:
        return list(self.tables)


class SqlAlchemyDatabaseExportRepository:
    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    async def add_job(self, job: DatabaseExportJob) -> DatabaseExportJob:
        self._session.add(job)
        await self._session.commit()
        await self._session.refresh(job)
        return job

    async def get_job(self, job_id: UUID) -> DatabaseExportJob | None:
        return await self._session.get(DatabaseExportJob, job_id)

    async def save_job(self, job: DatabaseExportJob) -> DatabaseExportJob:
        await self._session.commit()
        await self._session.refresh(job)
        return job

    async def load_tables(self) -> list[DatabaseExportTable]:
        tables: list[DatabaseExportTable] = []
        for table in sorted(Base.metadata.sorted_tables, key=lambda item: item.name):
            tables.append(await self._load_table(table))
        return tables

    async def _load_table(self, table: Table) -> DatabaseExportTable:
        result = await self._session.execute(select(table))
        rows = [
            {
                column.name: DatabaseExportService.json_value(row[column.name])
                for column in table.columns
            }
            for row in result.mappings().all()
        ]
        return DatabaseExportTable(
            name=table.name,
            columns=[column.name for column in table.columns],
            rows=rows,
            foreign_keys=[
                DatabaseExportForeignKey(
                    table=table.name,
                    column=foreign_key.parent.name,
                    referred_table=foreign_key.column.table.name,
                    referred_column=foreign_key.column.name,
                )
                for foreign_key in sorted(
                    table.foreign_keys,
                    key=lambda item: (item.parent.name, item.column.table.name),
                )
            ],
        )


class DatabaseExportService:
    def __init__(
        self,
        *,
        repository: DatabaseExportRepository,
        storage: StorageBackend,
        dispatcher: DatabaseExportTaskDispatcher | None = None,
        now_provider: Callable[[], datetime] = lambda: datetime.now(UTC),
    ) -> None:
        self._repository = repository
        self._storage = storage
        self._dispatcher = dispatcher or NoopDatabaseExportTaskDispatcher()
        self._now_provider = now_provider

    async def create_export(
        self,
        *,
        actor: User,
        audit_writer: AuditLogWriter | None = None,
        audit_context: AuditContext | None = None,
    ) -> DatabaseExportJob:
        self._ensure_admin(actor)
        now = self._now_provider()
        job = DatabaseExportJob(
            id=uuid4(),
            requested_by_id=actor.id,
            status=DatabaseExportJobStatus.queued,
            progress=0,
            table_count=0,
            row_count=0,
            storage_key=None,
            manifest={},
            error_message=None,
            started_at=None,
            finished_at=None,
            created_at=now,
            updated_at=now,
        )
        await self._repository.add_job(job)
        self._dispatcher.enqueue(job.id)
        self._record_audit(
            action="database_export.create_job",
            actor=actor,
            job=job,
            audit_writer=audit_writer,
            audit_context=audit_context,
            extra={},
        )
        return job

    async def get_export_job(self, *, actor: User, job_id: UUID) -> DatabaseExportJob:
        self._ensure_admin(actor)
        return await self._get_existing_job(job_id)

    async def run_export_job(self, job_id: UUID) -> DatabaseExportJob:
        job = await self._get_existing_job(job_id)
        now = self._now_provider()
        job.status = DatabaseExportJobStatus.running
        job.progress = 10
        job.started_at = now
        job.error_message = None
        job.updated_at = now
        await self._repository.save_job(job)
        try:
            tables = await self._repository.load_tables()
            manifest = self._build_manifest(job=job, tables=tables, generated_at=now)
            package = self._render_zip_package(
                tables=tables,
                manifest=manifest,
                foreign_keys=self._collect_foreign_keys(tables),
            )
            job.storage_key = self._storage.save(
                sub_id="exports",
                phase_id=str(job.id),
                filename="database_export.zip",
                content=package,
            )
            job.status = DatabaseExportJobStatus.completed
            job.progress = 100
            job.table_count = len(tables)
            job.row_count = cast(int, manifest["total_rows"])
            job.manifest = manifest
            job.finished_at = self._now_provider()
            job.updated_at = job.finished_at
            return await self._repository.save_job(job)
        except Exception as exc:
            job.status = DatabaseExportJobStatus.failed
            job.progress = 100
            job.error_message = str(exc)
            job.finished_at = self._now_provider()
            job.updated_at = job.finished_at
            await self._repository.save_job(job)
            raise

    async def download_export(
        self,
        *,
        actor: User,
        job_id: UUID,
        audit_writer: AuditLogWriter | None = None,
        audit_context: AuditContext | None = None,
    ) -> DatabaseExportDownload:
        self._ensure_admin(actor)
        job = await self._get_existing_job(job_id)
        if job.status != DatabaseExportJobStatus.completed:
            raise BusinessException(
                code=3032,
                message="数据库导出任务尚未完成",
                status_code=409,
                data={"status": job.status.value},
            )
        if job.storage_key is None:
            raise ResourceNotFoundError("数据库导出文件不存在")
        self._record_audit(
            action="database_export.download",
            actor=actor,
            job=job,
            audit_writer=audit_writer,
            audit_context=audit_context,
            extra={"storage_key": job.storage_key},
        )
        return DatabaseExportDownload(
            file_name="database_export.zip",
            content_type="application/zip",
            content=self._storage.read(job.storage_key),
        )

    def _build_manifest(
        self,
        *,
        job: DatabaseExportJob,
        tables: Sequence[DatabaseExportTable],
        generated_at: datetime,
    ) -> dict[str, object]:
        table_items = [
            {
                "name": table.name,
                "row_count": len(table.rows),
                "columns": table.columns,
            }
            for table in tables
        ]
        return {
            "job_id": str(job.id),
            "generated_at": generated_at.isoformat(),
            "operator_id": str(job.requested_by_id),
            "table_count": len(table_items),
            "total_rows": sum(len(table.rows) for table in tables),
            "tables": table_items,
        }

    def _render_zip_package(
        self,
        *,
        tables: Sequence[DatabaseExportTable],
        manifest: dict[str, object],
        foreign_keys: Sequence[DatabaseExportForeignKey],
    ) -> bytes:
        audit_info = {
            "job_id": manifest["job_id"],
            "generated_at": manifest["generated_at"],
            "operator_id": manifest["operator_id"],
            "table_count": manifest["table_count"],
            "total_rows": manifest["total_rows"],
        }
        buffer = BytesIO()
        with ZipFile(buffer, mode="w", compression=ZIP_DEFLATED) as archive:
            archive.writestr("manifest.json", self._json_bytes(manifest))
            archive.writestr(
                "foreign_keys.json",
                self._json_bytes([foreign_key.__dict__ for foreign_key in foreign_keys]),
            )
            archive.writestr("audit.json", self._json_bytes(audit_info))
            for table in tables:
                archive.writestr(f"csv/{table.name}.csv", self._render_csv(table))
            archive.writestr("xlsx/database_export.xlsx", self._render_xlsx(tables))
        return buffer.getvalue()

    def _render_csv(self, table: DatabaseExportTable) -> bytes:
        output = StringIO()
        writer = csv.DictWriter(output, fieldnames=table.columns, lineterminator="\n")
        writer.writeheader()
        for row in table.rows:
            writer.writerow({column: self.csv_value(row.get(column)) for column in table.columns})
        return output.getvalue().encode("utf-8-sig")

    def _render_xlsx(self, tables: Sequence[DatabaseExportTable]) -> bytes:
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        used_titles: set[str] = set()
        for table in tables:
            worksheet = workbook.create_sheet(self._sheet_title(table.name, used_titles))
            worksheet.append(table.columns)
            for row in table.rows:
                worksheet.append([self.csv_value(row.get(column)) for column in table.columns])
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _sheet_title(table_name: str, used_titles: set[str]) -> str:
        base = table_name[:31] or "table"
        title = base
        suffix = 1
        while title in used_titles:
            suffix += 1
            title = f"{base[: 31 - len(str(suffix)) - 1]}_{suffix}"
        used_titles.add(title)
        return title

    @staticmethod
    def _collect_foreign_keys(
        tables: Sequence[DatabaseExportTable],
    ) -> list[DatabaseExportForeignKey]:
        return [
            foreign_key
            for table in tables
            for foreign_key in sorted(
                table.foreign_keys,
                key=lambda item: (item.table, item.column, item.referred_table),
            )
        ]

    async def _get_existing_job(self, job_id: UUID) -> DatabaseExportJob:
        job = await self._repository.get_job(job_id)
        if job is None:
            raise ResourceNotFoundError("数据库导出任务不存在")
        return job

    def _record_audit(
        self,
        *,
        action: str,
        actor: User,
        job: DatabaseExportJob,
        audit_writer: AuditLogWriter | None,
        audit_context: AuditContext | None,
        extra: dict[str, object],
    ) -> None:
        if audit_writer is None:
            return
        context = audit_context or AuditContext(actor_id=actor.id)
        audit_writer.enqueue(
            AuditLogEntry(
                actor_id=context.actor_id or actor.id,
                action=action,
                target_type="database_export_job",
                target_id=str(job.id),
                before_state={},
                after_state={
                    "status": job.status.value,
                    "progress": job.progress,
                    "table_count": job.table_count,
                    "row_count": job.row_count,
                    "storage_key": job.storage_key,
                },
                ip_address=context.ip_address,
                user_agent=context.user_agent,
                extra=extra,
                request_id=context.request_id,
            ),
        )

    @staticmethod
    def _ensure_admin(actor: User) -> None:
        if actor.role != UserRole.admin:
            raise PermissionDeniedError()

    @staticmethod
    def _json_bytes(value: object) -> bytes:
        return json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8")

    @staticmethod
    def csv_value(value: object | None) -> object:
        if value is None:
            return ""
        if isinstance(value, DatabaseExportJobStatus):
            return value.value
        if isinstance(value, Decimal):
            return str(value)
        if isinstance(value, date | datetime):
            return value.isoformat()
        if isinstance(value, UUID):
            return str(value)
        if isinstance(value, dict | list | tuple):
            return json.dumps(DatabaseExportService.json_value(value), ensure_ascii=False)
        return value

    @staticmethod
    def json_value(value: object | None) -> object:
        if value is None:
            return None
        if isinstance(value, DatabaseExportJobStatus):
            return value.value
        if isinstance(value, Decimal | UUID):
            return str(value)
        if isinstance(value, date | datetime):
            return value.isoformat()
        if isinstance(value, dict):
            return {str(key): DatabaseExportService.json_value(item) for key, item in value.items()}
        if isinstance(value, list | tuple):
            return [DatabaseExportService.json_value(item) for item in value]
        return cast(Any, value)
