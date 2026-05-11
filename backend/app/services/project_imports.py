from __future__ import annotations

from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from secrets import token_hex
from time import perf_counter
from typing import Protocol
from uuid import UUID, uuid4

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import PermissionDeniedError, ValidationFailedError
from app.models.departments import Department
from app.models.main_projects import MainProject, MainProjectStatus
from app.models.users import User, UserRole
from app.services.audit import AuditContext, AuditLogEntry, AuditLogWriter

PROJECT_IMPORT_TEMPLATE_HEADERS = [
    "项目编号",
    "项目名称",
    "部门编码",
    "总预算",
    "预计完成日期",
    "备注",
]


@dataclass(frozen=True)
class ProjectImportColumn:
    key: str
    header: str
    width: int


PROJECT_IMPORT_COLUMNS = (
    ProjectImportColumn("project_no", "项目编号", 18),
    ProjectImportColumn("name", "项目名称", 28),
    ProjectImportColumn("dept_code", "部门编码", 16),
    ProjectImportColumn("total_budget", "总预算", 16),
    ProjectImportColumn("expected_finish_date", "预计完成日期", 18),
    ProjectImportColumn("remark", "备注", 30),
)


@dataclass(frozen=True)
class ProjectImportRowError:
    row_number: int
    field: str
    message: str
    value: str | None = None


@dataclass(frozen=True)
class ProjectImportCreatedProject:
    id: UUID
    project_no: str
    name: str
    dept_id: UUID
    status: MainProjectStatus


@dataclass(frozen=True)
class ProjectImportResult:
    batch_no: str
    total_rows: int
    success_count: int
    failure_count: int
    duration_ms: int
    errors: list[ProjectImportRowError]
    created_projects: list[ProjectImportCreatedProject]


@dataclass(frozen=True)
class ProjectImportRawRow:
    row_number: int
    project_no: object | None
    name: object | None
    dept_code: object | None
    total_budget: object | None
    expected_finish_date: object | None
    remark: object | None


@dataclass(frozen=True)
class ProjectImportValidatedRow:
    row_number: int
    project_no: str
    name: str
    department: Department
    total_budget: Decimal
    expected_finish_date: date
    remark: str | None


class ProjectImportRepository(Protocol):
    async def list_existing_project_nos(self, project_nos: Sequence[str]) -> set[str]:
        ...

    async def list_departments_by_codes(self, codes: Sequence[str]) -> dict[str, Department]:
        ...

    def add_project(self, project: MainProject) -> None:
        ...

    async def commit(self) -> None:
        ...

    async def refresh_projects(self, projects: Sequence[MainProject]) -> None:
        ...


class SqlAlchemyProjectImportRepository:
    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    async def list_existing_project_nos(self, project_nos: Sequence[str]) -> set[str]:
        if not project_nos:
            return set()
        result = await self._session.scalars(
            select(MainProject.project_no).where(MainProject.project_no.in_(project_nos)),
        )
        return set(result.all())

    async def list_departments_by_codes(self, codes: Sequence[str]) -> dict[str, Department]:
        if not codes:
            return {}
        result = await self._session.scalars(select(Department).where(Department.code.in_(codes)))
        return {department.code: department for department in result.all()}

    def add_project(self, project: MainProject) -> None:
        self._session.add(project)

    async def commit(self) -> None:
        await self._session.commit()

    async def refresh_projects(self, projects: Sequence[MainProject]) -> None:
        for project in projects:
            await self._session.refresh(project)


class InMemoryProjectImportRepository:
    def __init__(
        self,
        *,
        departments: Sequence[Department],
        projects: Sequence[MainProject] | None = None,
    ) -> None:
        self.departments = list(departments)
        self.projects = list(projects or [])

    async def list_existing_project_nos(self, project_nos: Sequence[str]) -> set[str]:
        requested = set(project_nos)
        return {project.project_no for project in self.projects if project.project_no in requested}

    async def list_departments_by_codes(self, codes: Sequence[str]) -> dict[str, Department]:
        requested = set(codes)
        return {
            department.code: department
            for department in self.departments
            if department.code in requested
        }

    def add_project(self, project: MainProject) -> None:
        self.projects.append(project)

    async def commit(self) -> None:
        return None

    async def refresh_projects(self, projects: Sequence[MainProject]) -> None:
        _ = projects
        return None


class ProjectImportService:
    def __init__(
        self,
        *,
        repository: ProjectImportRepository,
        now_provider: Callable[[], datetime] | None = None,
        batch_token_provider: Callable[[], str] | None = None,
    ) -> None:
        self._repository = repository
        self._now_provider = now_provider or _utc_now
        self._batch_token_provider = batch_token_provider or _new_batch_token

    def generate_template(self) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "项目导入模板"
        worksheet.append(PROJECT_IMPORT_TEMPLATE_HEADERS)
        worksheet.freeze_panes = "A2"
        for index, column in enumerate(PROJECT_IMPORT_COLUMNS, start=1):
            worksheet.column_dimensions[worksheet.cell(row=1, column=index).column_letter].width = (
                column.width
            )
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    async def import_projects(
        self,
        *,
        actor: User,
        content: bytes,
        audit_writer: AuditLogWriter | None = None,
        audit_context: AuditContext | None = None,
    ) -> ProjectImportResult:
        if actor.role != UserRole.admin:
            raise PermissionDeniedError()

        started = perf_counter()
        now = self._now_provider()
        batch_no = self._build_batch_no(now)
        raw_rows = self._load_rows(content)
        existing_project_nos = await self._repository.list_existing_project_nos(
            [project_no for project_no in self._extract_string_values(raw_rows, "project_no")],
        )
        departments_by_code = await self._repository.list_departments_by_codes(
            [dept_code for dept_code in self._extract_string_values(raw_rows, "dept_code")],
        )

        validated_rows: list[ProjectImportValidatedRow] = []
        errors: list[ProjectImportRowError] = []
        failed_rows: set[int] = set()
        seen_project_nos: set[str] = set()
        for raw_row in raw_rows:
            row, row_errors = self._validate_row(
                raw_row=raw_row,
                existing_project_nos=existing_project_nos,
                seen_project_nos=seen_project_nos,
                departments_by_code=departments_by_code,
            )
            if row_errors:
                errors.extend(row_errors)
                failed_rows.add(raw_row.row_number)
                continue
            if row is not None:
                validated_rows.append(row)
                seen_project_nos.add(row.project_no)

        created_projects = self._create_projects(
            rows=validated_rows,
            actor=actor,
            now=now,
        )
        for project in created_projects:
            self._repository.add_project(project)
        if created_projects:
            await self._repository.commit()
            await self._repository.refresh_projects(created_projects)

        result = ProjectImportResult(
            batch_no=batch_no,
            total_rows=len(raw_rows),
            success_count=len(created_projects),
            failure_count=len(failed_rows),
            duration_ms=int((perf_counter() - started) * 1000),
            errors=errors,
            created_projects=[
                ProjectImportCreatedProject(
                    id=project.id,
                    project_no=project.project_no,
                    name=project.name,
                    dept_id=project.dept_id,
                    status=project.status,
                )
                for project in created_projects
            ],
        )
        self._record_audit(
            actor=actor,
            result=result,
            audit_writer=audit_writer,
            audit_context=audit_context,
        )
        return result

    def _load_rows(self, content: bytes) -> list[ProjectImportRawRow]:
        try:
            workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        except Exception as exc:  # pragma: no cover - openpyxl exception types are broad
            raise ValidationFailedError(
                message="无法读取 Excel 文件",
                data={"reason": str(exc)},
            ) from exc

        try:
            worksheet = workbook.active
            header_row = next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
                (),
            )
            header_indexes = self._resolve_header_indexes(header_row)
            rows: list[ProjectImportRawRow] = []
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                if self._is_blank_row(values):
                    continue
                rows.append(
                    ProjectImportRawRow(
                        row_number=row_number,
                        project_no=self._value_at(values, header_indexes["project_no"]),
                        name=self._value_at(values, header_indexes["name"]),
                        dept_code=self._value_at(values, header_indexes["dept_code"]),
                        total_budget=self._value_at(values, header_indexes["total_budget"]),
                        expected_finish_date=self._value_at(
                            values,
                            header_indexes["expected_finish_date"],
                        ),
                        remark=self._value_at(values, header_indexes["remark"]),
                    ),
                )
            return rows
        finally:
            workbook.close()

    def _resolve_header_indexes(self, header_row: Sequence[object | None]) -> dict[str, int]:
        normalized_headers = [self._normalize_text(value) or "" for value in header_row]
        missing_headers = [
            column.header
            for column in PROJECT_IMPORT_COLUMNS
            if column.header not in normalized_headers
        ]
        if missing_headers:
            raise ValidationFailedError(
                message="导入模板字段不完整",
                data={"missing_headers": missing_headers},
            )
        return {
            column.key: normalized_headers.index(column.header)
            for column in PROJECT_IMPORT_COLUMNS
        }

    def _validate_row(
        self,
        *,
        raw_row: ProjectImportRawRow,
        existing_project_nos: set[str],
        seen_project_nos: set[str],
        departments_by_code: dict[str, Department],
    ) -> tuple[ProjectImportValidatedRow | None, list[ProjectImportRowError]]:
        errors: list[ProjectImportRowError] = []
        project_no = self._normalize_required_text(
            raw_row.project_no,
            field="project_no",
            label="项目编号",
            row_number=raw_row.row_number,
            errors=errors,
            max_length=32,
        )
        name = self._normalize_required_text(
            raw_row.name,
            field="name",
            label="项目名称",
            row_number=raw_row.row_number,
            errors=errors,
            max_length=200,
        )
        dept_code = self._normalize_required_text(
            raw_row.dept_code,
            field="dept_code",
            label="部门编码",
            row_number=raw_row.row_number,
            errors=errors,
            max_length=50,
        )
        total_budget = self._parse_budget(raw_row)
        if total_budget is None:
            errors.append(
                ProjectImportRowError(
                    row_number=raw_row.row_number,
                    field="total_budget",
                    message="总预算必须是非负金额",
                    value=self._stringify_value(raw_row.total_budget),
                ),
            )
        expected_finish_date = self._parse_date(raw_row.expected_finish_date)
        if expected_finish_date is None:
            errors.append(
                ProjectImportRowError(
                    row_number=raw_row.row_number,
                    field="expected_finish_date",
                    message="预计完成日期必须是日期或 YYYY-MM-DD 文本",
                    value=self._stringify_value(raw_row.expected_finish_date),
                ),
            )

        department = departments_by_code.get(dept_code or "")
        if dept_code and department is None:
            errors.append(
                ProjectImportRowError(
                    row_number=raw_row.row_number,
                    field="dept_code",
                    message="部门编码不存在",
                    value=dept_code,
                ),
            )
        if project_no and (project_no in existing_project_nos or project_no in seen_project_nos):
            errors.append(
                ProjectImportRowError(
                    row_number=raw_row.row_number,
                    field="project_no",
                    message="项目编号重复",
                    value=project_no,
                ),
            )

        if (
            errors
            or project_no is None
            or name is None
            or department is None
            or total_budget is None
            or expected_finish_date is None
        ):
            return None, errors

        return (
            ProjectImportValidatedRow(
                row_number=raw_row.row_number,
                project_no=project_no,
                name=name,
                department=department,
                total_budget=total_budget,
                expected_finish_date=expected_finish_date,
                remark=self._normalize_text(raw_row.remark),
            ),
            [],
        )

    def _create_projects(
        self,
        *,
        rows: Sequence[ProjectImportValidatedRow],
        actor: User,
        now: datetime,
    ) -> list[MainProject]:
        return [
            MainProject(
                id=uuid4(),
                project_no=row.project_no,
                name=row.name,
                dept_id=row.department.id,
                status=MainProjectStatus.not_started,
                total_budget=row.total_budget,
                expected_finish_date=row.expected_finish_date,
                spent_amount=Decimal("0.00"),
                remark=row.remark,
                creator_id=actor.id,
                created_at=now,
                updated_at=now,
            )
            for row in rows
        ]

    def _build_batch_no(self, now: datetime) -> str:
        return f"IMPORT-{now:%Y%m%d}-{self._batch_token_provider()}"

    def _record_audit(
        self,
        *,
        actor: User,
        result: ProjectImportResult,
        audit_writer: AuditLogWriter | None,
        audit_context: AuditContext | None,
    ) -> None:
        if audit_writer is None:
            return
        context = audit_context or AuditContext(actor_id=actor.id)
        audit_writer.enqueue(
            AuditLogEntry(
                actor_id=context.actor_id or actor.id,
                action="project_import.create_batch",
                target_type="project_import_batch",
                target_id=result.batch_no,
                before_state={},
                after_state={
                    "batch_no": result.batch_no,
                    "total_rows": result.total_rows,
                    "success_count": result.success_count,
                    "failure_count": result.failure_count,
                    "duration_ms": result.duration_ms,
                },
                ip_address=context.ip_address,
                user_agent=context.user_agent,
                extra={
                    "created_project_nos": [
                        project.project_no for project in result.created_projects
                    ],
                    "error_count": len(result.errors),
                },
                request_id=context.request_id,
            ),
        )

    def _extract_string_values(
        self,
        rows: Sequence[ProjectImportRawRow],
        field_name: str,
    ) -> list[str]:
        values: list[str] = []
        for row in rows:
            value = self._normalize_text(getattr(row, field_name))
            if value is not None:
                values.append(value)
        return values

    def _normalize_required_text(
        self,
        value: object | None,
        *,
        field: str,
        label: str,
        row_number: int,
        errors: list[ProjectImportRowError],
        max_length: int,
    ) -> str | None:
        normalized = self._normalize_text(value)
        if normalized is None:
            errors.append(
                ProjectImportRowError(
                    row_number=row_number,
                    field=field,
                    message=f"{label}不能为空",
                    value=None,
                ),
            )
            return None
        if len(normalized) > max_length:
            errors.append(
                ProjectImportRowError(
                    row_number=row_number,
                    field=field,
                    message=f"{label}长度不能超过 {max_length}",
                    value=normalized,
                ),
            )
            return None
        return normalized

    def _parse_budget(self, raw_row: ProjectImportRawRow) -> Decimal | None:
        value = raw_row.total_budget
        if value is None:
            return None
        try:
            amount = Decimal(str(value).strip()).quantize(Decimal("0.01"))
        except (InvalidOperation, ValueError):
            return None
        if amount < Decimal("0.00"):
            return None
        if len(amount.as_tuple().digits) > 15:
            return None
        return amount

    def _parse_date(self, value: object | None) -> date | None:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = self._normalize_text(value)
        if text is None:
            return None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _normalize_text(value: object | None) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    @staticmethod
    def _stringify_value(value: object | None) -> str | None:
        if value is None:
            return None
        return str(value)

    @staticmethod
    def _is_blank_row(values: Sequence[object | None]) -> bool:
        return all(value is None or str(value).strip() == "" for value in values)

    @staticmethod
    def _value_at(values: Sequence[object | None], index: int) -> object | None:
        if index >= len(values):
            return None
        return values[index]


def _utc_now() -> datetime:
    return datetime.now(UTC)


def _new_batch_token() -> str:
    return token_hex(4).upper()
